"""输入解析层：把邮件正文与多格式附件统一为结构化块。

    块类型：
      {"type": "text",  "text": "..."}                              文本块
      {"type": "table", "header": [...], "rows": [[...], ...]}     表格块
    flags: ["scanned"] 表示图片型 PDF（无文字层），走降级路径。
    """
import pathlib

from openpyxl import load_workbook
import pdfplumber
from docx import Document as DocxDocument


def parse_body(path):
    return {"name": pathlib.Path(path).name, "kind": "body",
            "blocks": [{"type": "text", "text": pathlib.Path(path).read_text(encoding="utf-8")}]}


def parse_xlsx(path):
    wb = load_workbook(path, data_only=True)
    blocks = []
    for ws in wb.worksheets:
        matrix = [[("" if c is None else str(c).strip()) for c in row] for row in ws.iter_rows(values_only=True)]
        # 矩阵尾部可能带大量空行，先裁掉
        while matrix and all(not c for c in matrix[-1]):
            matrix.pop()
        if not matrix:
            continue
        # 首个含 >=3 个非空单元格的行视为表头
        header_idx = None
        for i, row in enumerate(matrix):
            if sum(1 for c in row if c) >= 3:
                header_idx = i
                break
        if header_idx is None:
            text = "\n".join(" ".join(c for c in row if c) for row in matrix)
            blocks.append({"type": "text", "text": text})
            continue
        header = matrix[header_idx]
        width = len(header)
        rows = [r[:width] + [""] * (width - len(r)) for r in matrix[header_idx + 1:]]
        rows = [r for r in rows if any(c for c in r)]
        blocks.append({"type": "table", "header": header, "rows": rows})
    return {"name": pathlib.Path(path).name, "kind": "xlsx", "blocks": blocks}


def parse_docx(path):
    doc = DocxDocument(path)
    blocks = []
    for p in doc.paragraphs:
        if p.text.strip():
            blocks.append({"type": "text", "text": p.text.strip()})
    for t in doc.tables:
        header = [c.text.strip() for c in t.rows[0].cells]
        rows = [[c.text.strip() for c in r.cells] for r in t.rows[1:]]
        if any(header):
            blocks.append({"type": "table", "header": header, "rows": rows})
    return {"name": pathlib.Path(path).name, "kind": "docx", "blocks": blocks}


def parse_pdf(path):
    flags = []
    pages_text = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages_text.append(page.extract_text() or "")
    total = "".join(pages_text)
    if len(total.strip()) < 20:
        flags.append("scanned")
    blocks = [{"type": "text", "text": t} for t in pages_text if t.strip()]
    return {"name": pathlib.Path(path).name, "kind": "pdf", "blocks": blocks, "flags": flags}


def parse_attachment(path):
    p = pathlib.Path(path)
    suffix = p.suffix.lower()
    if suffix == ".xlsx":
        return parse_xlsx(p)
    if suffix == ".docx":
        return parse_docx(p)
    if suffix == ".pdf":
        return parse_pdf(p)
    if suffix in (".txt", ".csv"):
        return {"name": p.name, "kind": "txt",
                "blocks": [{"type": "text", "text": p.read_text(encoding="utf-8")}]}
    raise ValueError(f"不支持的附件格式：{suffix}")


def try_ocr(att):
    """OCR 兜底钩子：默认返回 None（未接入引擎）。

    真实环境可在此接入 pytesseract / paddleocr，返回 {"text": ...}。
    """
    return None
