"""导出层：把复核后的记录写成带标色的标准报岗表 xlsx。"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from .schema import FIELDS, FIELD_LABELS, DERIVED_FIELDS, ASSUMPTIONS

FILLS = {
    "green": "EAF3DE",
    "amber": "FAEEDA",
    "red": "FCEBEB",
    "edited": "E6F1FB",
}


def export_xlsx(template_name, records, mode):
    wb = Workbook()
    ws = wb.active
    ws.title = "报岗表"

    headers = [f["label"] for f in FIELDS] + [DERIVED_FIELDS[0]["label"], "置信状态", "复核说明"]
    ws.append([f"{template_name} · 入离职报岗表（抽取模式：{mode}）"])
    ws.append(headers)

    for rec in records:
        row = [rec["values"].get(f["key"]) or "" for f in FIELDS]
        row.append(rec.get("derived", {}).get("social_month", ""))
        edited_any = any(rec["fields"][f["key"]]["edited"] for f in FIELDS)
        if edited_any:
            row.append("含人工确认字段")
        elif rec["status"] == "ok":
            row.append("全部正常")
        else:
            row.append("待复核")
        row.append("；".join(i["message"] for i in rec["issues"]) or "—")
        ws.append(row)
        r = ws.max_row
        for ci, f in enumerate(FIELDS, start=1):
            cell = ws.cell(row=r, column=ci)
            if rec["fields"][f["key"]]["edited"]:
                cell.fill = PatternFill("solid", start_color=FILLS["edited"])
            else:
                level = rec["fields"][f["key"]]["level"]
                if level in FILLS:
                    cell.fill = PatternFill("solid", start_color=FILLS[level])

    ws2 = wb.create_sheet("复核清单")
    ws2.append(["记录", "姓名", "字段", "规则", "级别", "说明"])
    for idx, rec in enumerate(records):
        name = rec["values"].get("name") or rec.get("key", "")
        for i in rec["issues"]:
            ws2.append([idx + 1, name, FIELD_LABELS.get(i.get("field"), i.get("field") or "—"),
                        i.get("rule", ""), i.get("level", ""), i.get("message", "")])
    if not any(len(r) > 1 and r[2] for r in ws2.iter_rows(min_row=2, values_only=True)):
        ws2.append([1, "—", "—", "—", "—", "无待复核项"])

    ws3 = wb.create_sheet("口径说明")
    ws3.append(["本表由转置小助手生成，全部数据为合成样本，非真实个人信息"])
    ws3.append([])
    ws3.append(["口径假设（须在真实对接阶段由业务方确认后替换）："])
    for n in ASSUMPTIONS["notes"]:
        ws3.append([n])
    ws3.append([f"社保基数区间：{ASSUMPTIONS['base_range'][0]} – {ASSUMPTIONS['base_range'][1]}"])
    ws3.append([f"生效日期窗口：{ASSUMPTIONS['effective_date_window'][0]} – {ASSUMPTIONS['effective_date_window'][1]}"])
    ws3.append([f"最低用工年龄：{ASSUMPTIONS['min_employment_age']} 周岁"])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
