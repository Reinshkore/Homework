"""生成 7 份合成邮件样本（3 客户模板 × 2 封 + 1 扫描件）。

    所有姓名、证件号均为编造；证件号使用校验位正确的虚构号码，
    避免与真实身份碰撞。运行：python scripts/gen_samples.py
    """
import json
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

SAMPLES = ROOT / "samples"

ID_WEIGHTS = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
ID_CHECK_MAP = "10X98765432"


def make_id(region, birth, seq):
    body = f"{region}{birth}{seq}"
    s = sum(int(body[i]) * ID_WEIGHTS[i] for i in range(17))
    return body + ID_CHECK_MAP[s % 11]


def corrupt(idv):
    last = idv[-1]
    return idv[:-1] + ("X" if last != "X" else "0")


def find_font():
    for name in ("simhei.ttf", "simfang.ttf", "msyh.ttc"):
        p = pathlib.Path("C:/Windows/Fonts") / name
        if p.exists():
            return p
    raise RuntimeError("未找到可用中文字体（simhei.ttf / simfang.ttf）")


FONT_PATH = find_font()
pdfmetrics.registerFont(TTFont("CNFONT", str(FONT_PATH)))


def write_body(d, text):
    (SAMPLES / d).mkdir(parents=True, exist_ok=True)
    (SAMPLES / d / "body.txt").write_text(text, encoding="utf-8")


def make_xlsx(d, fname, header, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(SAMPLES / d / fname)


def make_docx(d, fname, paragraphs):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(SAMPLES / d / fname)


def make_pdf(d, fname, lines):
    c = pdfcanvas.Canvas(str(SAMPLES / d / fname), pagesize=(595, 842))
    y = 800
    for line in lines:
        c.setFont("CNFONT", 12)
        c.drawString(50, y, line)
        y -= 24
        if y < 60:
            c.showPage()
            y = 800
    c.save()


def make_scanned_pdf(d, fname, lines):
    font = ImageFont.truetype(str(FONT_PATH), 22)
    img = Image.new("RGB", (780, 60 + 40 * len(lines)), "white")
    draw = ImageDraw.Draw(img)
    y = 24
    for line in lines:
        draw.text((28, y), line, fill="black", font=font)
        y += 40
    tmp = SAMPLES / d / "_scan_tmp.png"
    img.save(tmp)
    c = pdfcanvas.Canvas(str(SAMPLES / d / fname), pagesize=(595, 842))
    c.drawImage(str(tmp), 40, 500, width=515, height=515 * img.height / img.width)
    c.save()
    tmp.unlink()


def main():
    ids = {
        "张明": make_id("310101", "19900307", "001"),
        "李华": make_id("110105", "19871120", "002"),
        "王芳": make_id("440304", "19950601", "003"),
        "刘洋": make_id("310101", "19920315", "451"),
        "陈静": make_id("110105", "19950822", "342"),
        "赵磊": make_id("440304", "19910718", "523"),
        "周杰": make_id("310101", "19910823", "401"),
        "吴敏": make_id("110105", "19920615", "362"),
        "郑浩": make_id("310101", "19890412", "673"),
        "何娜": make_id("110105", "19940709", "824"),
        "王强": make_id("310101", "19880512", "123"),
        "孙丽": make_id("440304", "19921207", "564"),
        "郑凯": make_id("310101", "19900111", "783"),
        "冯婷": make_id("110105", "20120506", "234"),
        "钱进": make_id("310101", "19930718", "115"),
        "许婷": make_id("110105", "19951123", "446"),
    }

    header_a = ["员工姓名", "证件号码", "职位", "工作城市", "社保缴纳基数", "类型", "生效日期"]

    # A-01：3 条正常记录
    write_body("A-01", "各位好：\n\n附件为本月入离职报岗名单，请查收并处理。\n\n上海外服 业务组")
    make_xlsx("A-01", "名单.xlsx", header_a, [
        ["张明", ids["张明"], "工程师", "上海", 8000, "入职", "2026/03/16"],
        ["李华", ids["李华"], "会计", "北京", 9000, "入职", "2026/03/16"],
        ["王芳", ids["王芳"], "行政专员", "深圳", 7000, "离职", "2026/03/31"],
    ])

    # A-02：校验位错 / 日期 2 月 30 日 / 必填缺失 + 枚举非法
    write_body("A-02", "各位好：\n\n附件名单中有一处信息待客户确认，请留意校验结果。\n\n上海外服 业务组")
    make_xlsx("A-02", "名单.xlsx", header_a, [
        ["刘洋", corrupt(ids["刘洋"]), "物流专员", "上海", 8200, "入职", "2026/03/20"],
        ["陈静", ids["陈静"], "人事专员", "北京", 8500, "入职", "2026/02/30"],
        ["赵磊", ids["赵磊"], "", "深圳", 7800, "调岗", "2026/04/01"],
    ])

    # B-01：叙述型，2 条正常
    write_body("B-01", "您好，正文即名单，无附件表格，请直接处理。\n\n外服云 客户成功组")
    make_docx("B-01", "通知.docx", [
        "尊敬的外服同事：现将本月人员变动情况通知如下，请协助办理。",
        f"兹通知人员变动：周杰（身份证 {ids['周杰']}），岗位市场专员，工作地上海，"
        f"自 2026年4月1日 起入职，社保基数 8200。",
        f"兹通知人员变动：吴敏（身份证 {ids['吴敏']}），岗位客服主管，工作地北京，"
        f"自 2026年3月31日 起离职。",
    ])

    # B-02：日期 13 月（不可解析）/ 入职缺社保基数
    write_body("B-02", "您好，正文即名单，请直接处理，其中一人的信息需与客户再核实。\n\n外服云 客户成功组")
    make_docx("B-02", "通知.docx", [
        "尊敬的外服同事：现将本月人员变动情况通知如下，请协助办理。",
        f"兹通知人员变动：郑浩（身份证 {ids['郑浩']}），岗位销售，工作地上海，"
        f"自 2026年13月5日 起入职，社保基数 7500。",
        f"兹通知人员变动：何娜（身份证 {ids['何娜']}），岗位行政，工作地北京，"
        f"自 2026年4月10日 起入职。",
    ])

    # C-01：PDF 表单，2 条正常（离职记录无基数）
    write_body("C-01", "您好，附件为客户C本月人员变动通知，请查收。\n\n客户C 人力资源部")
    make_pdf("C-01", "通知.pdf", [
        "客户C人力资源部：",
        "现将本月人员变动通知如下。",
        "【员工一】",
        f"姓名：王强",
        f"证件号码：{ids['王强']}",
        "岗位：财务主管",
        "工作地：上海",
        "类型：入职",
        "生效日期：2026-04-01",
        "社保基数按 8000 执行",
        "",
        "【员工二】",
        f"姓名：孙丽",
        f"证件号码：{ids['孙丽']}",
        "岗位：采购专员",
        "工作地：深圳",
        "类型：离职",
        "生效日期：2026-04-15",
    ])

    # C-02：社保基数跨源冲突（附件 8000 vs 正文 8500）/ 未满 16 周岁
    write_body("C-02",
               "您好，附件为客户C本月人员变动通知。其中，郑凯的社保基数有更新，"
               "为 8500，请以本邮件为准。\n\n客户C 人力资源部")
    make_pdf("C-02", "通知.pdf", [
        "客户C人力资源部：",
        "现将本月人员变动通知如下。",
        "【员工一】",
        f"姓名：郑凯",
        f"证件号码：{ids['郑凯']}",
        "岗位：物流主管",
        "工作地：上海",
        "类型：入职",
        "生效日期：2026-04-01",
        "社保基数按 8000 执行",
        "",
        "【员工二】",
        f"姓名：冯婷",
        f"证件号码：{ids['冯婷']}",
        "岗位：实习生",
        "工作地：北京",
        "类型：入职",
        "生效日期：2026-04-01",
    ])

    # S-01：扫描件（无文字层），正文给出部分线索
    write_body("S-01",
               "各位好：\n\n附件为纸质名单扫描件，本周人员变动如下：\n\n"
               f"钱进（入职）\n\n许婷（离职）\n\n详情请见扫描附件，麻烦人工核对。")
    make_scanned_pdf("S-01", "扫描件.pdf", [
        "员工姓名    证件号码              职业类别    工作城市    社保缴纳基数    类型    生效日期",
        f"钱进        {ids['钱进']}    物流专员    上海        8100            入职    2026/03/23",
        f"许婷        {ids['许婷']}    行政助理    北京        7300            离职    2026/03/31",
    ])

    meta = [
        {"sample_id": "A-01", "template_id": "CUST_A", "dir": "A-01",
         "desc": "客户A · xlsx 表格 · 3 条正常记录",
         "files": {"body": "body.txt", "attachments": ["名单.xlsx"]}},
        {"sample_id": "A-02", "template_id": "CUST_A", "dir": "A-02",
         "desc": "客户A · xlsx · 埋错：证件校验位 / 日期非法 / 必填缺失 + 枚举非法",
         "files": {"body": "body.txt", "attachments": ["名单.xlsx"]}},
        {"sample_id": "B-01", "template_id": "CUST_B", "dir": "B-01",
         "desc": "客户B · docx 叙述 · 2 条正常记录",
         "files": {"body": "body.txt", "attachments": ["通知.docx"]}},
        {"sample_id": "B-02", "template_id": "CUST_B", "dir": "B-02",
         "desc": "客户B · docx · 埋错：生效日期 13 月 / 入职缺社保基数",
         "files": {"body": "body.txt", "attachments": ["通知.docx"]}},
        {"sample_id": "C-01", "template_id": "CUST_C", "dir": "C-01",
         "desc": "客户C · PDF 表单 · 2 条正常记录",
         "files": {"body": "body.txt", "attachments": ["通知.pdf"]}},
        {"sample_id": "C-02", "template_id": "CUST_C", "dir": "C-02",
         "desc": "客户C · PDF · 埋错：社保基数跨源冲突 / 未满 16 周岁",
         "files": {"body": "body.txt", "attachments": ["通知.pdf"]}},
        {"sample_id": "S-01", "template_id": "CUST_A", "dir": "S-01",
         "desc": "扫描件 · 无文字层 · 降级路径演示",
         "files": {"body": "body.txt", "attachments": ["扫描件.pdf"]}},
    ]
    (SAMPLES / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")

    # 生成一个"校验位正确性"的自检：确认生成的证件号都通过校验
    from core.validate import validate_id_no
    for n, v in ids.items():
        assert validate_id_no(v), f"{n} 的证件号校验失败：{v}"
    assert not validate_id_no(corrupt(ids["刘洋"])), "corrupt() 未产生非法校验位"
    print(f"OK：7 份样本已生成于 {SAMPLES}，{len(ids)} 个虚构证件号全部校验通过")


if __name__ == "__main__":
    main()
